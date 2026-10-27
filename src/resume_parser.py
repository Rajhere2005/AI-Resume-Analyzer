"""Extract text from PDF/DOCX resumes. Files are processed in memory only and never executed or saved."""
import io
import re

import fitz  # PyMuPDF
from docx import Document

MAX_FILE_MB = 5
MAX_PAGES = 10
MIN_CHARS = 100
MAX_CHARS = 20_000


class ResumeParseError(Exception):
    """Raised with a user-friendly message when a resume can't be read."""


def extract_text(filename: str, data: bytes) -> str:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext not in ("pdf", "docx"):
        raise ResumeParseError("Unsupported file type. Please upload a PDF or DOCX file.")
    if not data:
        raise ResumeParseError("The uploaded file is empty.")
    if len(data) > MAX_FILE_MB * 1024 * 1024:
        raise ResumeParseError(f"File is too large. The limit is {MAX_FILE_MB} MB.")

    # Check magic bytes so a renamed file can't masquerade as a resume.
    if ext == "pdf":
        if not data.startswith(b"%PDF"):
            raise ResumeParseError("This file doesn't look like a valid PDF.")
        text = _read_pdf(data)
    else:
        if not data.startswith(b"PK"):
            raise ResumeParseError("This file doesn't look like a valid DOCX document.")
        text = _read_docx(data)

    text = _clean(text)
    if len(text) < MIN_CHARS:
        raise ResumeParseError(
            "Very little text could be extracted. The file may be a scanned image or mostly graphics. "
            "Please upload a text-based PDF or DOCX (OCR is not supported)."
        )
    return text[:MAX_CHARS]


def _read_pdf(data: bytes) -> str:
    try:
        with fitz.open(stream=data, filetype="pdf") as doc:
            if doc.needs_pass:
                raise ResumeParseError("This PDF is password-protected. Please upload an unlocked copy.")
            if doc.page_count > MAX_PAGES:
                raise ResumeParseError(f"PDF has {doc.page_count} pages. Resumes are limited to {MAX_PAGES} pages.")
            return "\n".join(page.get_text("text") for page in doc)
    except ResumeParseError:
        raise
    except Exception:
        raise ResumeParseError("The PDF appears to be corrupted and couldn't be read.")


def _read_docx(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    except Exception:
        raise ResumeParseError("The DOCX file appears to be corrupted and couldn't be read.")


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
